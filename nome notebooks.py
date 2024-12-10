import os
from docx import Document
import pandas as pd

# Caminhos
pasta = r"C:\Users\Thiago_Mattos\Downloads\termos_colaboradores"  # Substitua pelo caminho da pasta dos arquivos Word
planilha_path = r"C:\Users\Thiago_Mattos\Downloads\Ativos de TI.xlsx"  # Substitua pelo caminho da planilha

# Carregar a planilha
df_planilha = pd.read_excel(planilha_path)

# Normalizar nomes para evitar problemas
df_planilha['Usuário/Responsável'] = df_planilha['Usuário/Responsável'].str.strip()

# Função para buscar configuração do notebook pelo primeiro e segundo nome
def buscar_configuracao_por_nomes(nomes, df):
    try:
        primeiro_nome, segundo_nome = nomes.split()[:2]
        condicao = df['Usuário/Responsável'].str.contains(f"{primeiro_nome}.*{segundo_nome}", case=False, na=False)
        dados = df[condicao].iloc[0]
        return dados.to_dict()
    except IndexError:
        return None

# Função para substituir placeholders e colocar em negrito
def substituir_placeholder_em_paragrafos(paragrafo, placeholder, texto, negrito=False):
    if placeholder in paragrafo.text:
        for run in paragrafo.runs:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, texto)
                if negrito:
                    run.bold = True

# Iterar pelos arquivos na pasta
for arquivo in os.listdir(pasta):
    if arquivo.endswith(".docx"):
        # Extrair o nome do funcionário do nome do arquivo
        nome_arquivo = arquivo.split("_", 1)[-1].replace(".docx", "").replace("_", " ")
        nome_funcionario = nome_arquivo.strip()
        
        # Buscar as configurações do notebook
        dados_notebook = buscar_configuracao_por_nomes(nome_funcionario, df_planilha)
        
        if dados_notebook:
            # Abrir o arquivo Word
            caminho_arquivo = os.path.join(pasta, arquivo)
            doc = Document(caminho_arquivo)
            
            # Substituir os placeholders no documento
            for paragrafo in doc.paragraphs:
                substituir_placeholder_em_paragrafos(paragrafo, "{{Modelo}}", dados_notebook["Modelo"], negrito=True)
                substituir_placeholder_em_paragrafos(paragrafo, "{{Processador}}", dados_notebook["Processador"], negrito=True)
                substituir_placeholder_em_paragrafos(paragrafo, "{{Ram}}", dados_notebook["Memória"], negrito=True)
                substituir_placeholder_em_paragrafos(paragrafo, "{{Armazenamento}}", dados_notebook["Armazenamento"], negrito=True)
                substituir_placeholder_em_paragrafos(paragrafo, "{{Identificador}}", dados_notebook["Identificador / SN"], negrito=True)
            
            # Salvar o arquivo atualizado
            doc.save(os.path.join(pasta, f"atualizado_{arquivo}"))
            print(f"Arquivo atualizado: {arquivo}")
        else:
            print(f"Configuração não encontrada para {nome_funcionario}")
